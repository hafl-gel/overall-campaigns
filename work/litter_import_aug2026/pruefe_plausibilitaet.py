from pathlib import Path
from datetime import date
import csv,json,math,statistics,hashlib
import openpyxl
base=Path('/Users/stefan/git/overall-campaigns')
out=base/'Einstreu';target=out/'litter_data_all.xlsx'
meta_path=Path('/Users/stefan/git/mettmenstetten-aug2026/data/Einstreuproben_Mettmenstetten.xlsx')
lab_path=Path('/Users/stefan/git/mettmenstetten-aug2026/data/HD_115914_HAFL  Gfeller Stefan_20260927.xlsx')
sha=hashlib.sha256(target.read_bytes()).hexdigest()
wb=openpyxl.load_workbook(target,data_only=True);s=wb['Sheet 1']
headers=[c.value for c in s[1]]
records=[dict(zip(headers,[c.value for c in row]),excel_row=row[0].row) for row in list(s)[1:]]
aug=[r for r in records if r['campaign']=='AUG26']
old=[r for r in records if r['campaign']!='AUG26' and r['stall'] in (1,2)]
meta={int(row[0].value):[c.value for c in row] for row in openpyxl.load_workbook(meta_path,data_only=True).active if isinstance(row[0].value,(int,float))}
labs={int(row[6].value):[c.value for c in row] for row in openpyxl.load_workbook(lab_path,data_only=True).active if isinstance(row[0].value,(int,float)) and row[5].value=='B61'}
assert len(aug)==31 and set(meta)==set(labs)=={r['Probenr-HAFL'] for r in aug}
assert len({(r['campaign'],r['stall'],r['masttag'],r['Wiederholung']) for r in aug})==31
mapping={'A':'F','B':'G','F':'H','G':'J','H':'P','I':'Q','J':'L','K':'M','L':'N','M':'R','N':'T','O':'U','P':'V','Q':'W','R':'X','S':'Y','T':'Z','U':'AA','V':'AC','X':'AD','Y':'AE','Z':'AF','AA':'AG'}
checks=[];n_source=0
for r in aug:
 nr=r['excel_row'];id=r['Probenr-HAFL'];lab=labs[id];m=meta[id]
 for dest,src in mapping.items():
  a=s[f'{dest}{nr}'].value;b=lab[openpyxl.utils.column_index_from_string(src)-1]
  assert a==b or isinstance(a,(float,int)) and isinstance(b,(float,int)) and abs(a-b)<1e-10,(id,dest,a,b)
  n_source+=1
 for dest,ix in [('C',2),('D',3),('E',4)]:
  assert s[f'{dest}{nr}'].value==m[ix];n_source+=1
 assert (m[1].date()-date(2026,7,24)).days+1==r['masttag']
 ts=r['TS_pct'];n=r['N_total_gkg_TS'];nh=r['NH4_gkg_TS']
 assert 0<ts<=100 and 0<=r['moisture_pct']<=100 and 0<=r['pH']<=14
 assert 0<=nh<=n and 0<=r['Kohlenstoff Corg g/kg TS']<=1000
 assert sum(r[k] for k in ['Calcium g/t TS','Kalium g/t TS','Magnesium g/t TS','Phosphor g/t TS','Schwefel g/t TS'])<=r['Glührückstand 500°C %']*10
 expressions=[
 ('Feuchte = 100 - TS',r['moisture_pct'],100-ts,1e-10),
 ('N_FS = N_TS * TS/100',r['N_total_gkg_FS'],n*ts/100,1e-10),
 ('Glühverlust + Asche = 100',r['Glühverlust 500°C (OS) % TS']+r['Glührückstand 500°C %'],100,.1000001),
 ('C/N = Corg/N',r['C/N-Verhältnis'],r['Kohlenstoff Corg g/kg TS']/n,.005+.05/n+.05*r['Kohlenstoff Corg g/kg TS']/n**2),
 ('N_FS Labor vs. aus angezeigten Werten',r['Stickstoff berechnet g/kg FS'],r['N_total_gkg_FS'],.005+.05*ts/100+.05*n/100+.000025),
 ('P2O5_TS = 2.291 * P_TS',r['Phosphorpentoxid g/t TS'],r['Phosphor g/t TS']*2.291,.005+2.291*.005),
 ('K2O_TS = 1.205 * K_TS',r['Kaliumoxid g/t TS'],r['Kalium g/t TS']*1.205,.005+1.205*.05),
 ('als Phosphor_FS beschrifteter Wert vs. P2O5_FS',r['Phosphor berechnet g/kg FS'],r['Phosphorpentoxid g/t TS']*ts/100,.005+.005*ts/100+.05*r['Phosphorpentoxid g/t TS']/100+.0000025)]
 for title,obs,exp,tol in expressions:
  checks.append(dict(sample_id=id,excel_row=nr,check=title,observed=obs,expected=exp,residual=obs-exp,tolerance=tol,within_rounding=abs(obs-exp)<=tol))
assert all(x['within_rounding'] for x in checks)
metrics=['pH','moisture_pct','N_total_gkg_TS','NH4_gkg_TS','C/N-Verhältnis']
groups=[]
for barn in (1,2):
 for day in (21,26,34):
  rr=[r for r in aug if r['stall']==barn and r['masttag']==day]
  g=dict(barn=barn,mastday=day,n=len(rr),sample_ids=','.join(str(r['Probenr-HAFL']) for r in rr))
  for k in metrics:
   a=[r[k] for r in rr];g[k+'_mean']=statistics.mean(a);g[k+'_min']=min(a);g[k+'_max']=max(a);g[k+'_sd']=statistics.stdev(a)
  groups.append(g)
ranges=[]
for camp in ['FEB','APR','JUL','DEC','AUG26']:
 rr=[r for r in records if r['campaign']==camp and r['stall'] in (1,2)]
 for k in metrics+['TS_pct']:
  vals=[r[k] for r in rr]
  ranges.append(dict(campaign=camp,metric=k,n=len(vals),minimum=min(vals),mean=statistics.mean(vals),median=statistics.median(vals),maximum=max(vals)))
flags=[]
floors={'pH':.3,'TS_pct':4,'N_total_gkg_TS':4,'NH4_gkg_TS':.8}
for barn in (1,2):
 for day in (21,26,34):
  rr=[r for r in aug if r['stall']==barn and r['masttag']==day]
  for k,minimum in floors.items():
   vals=[r[k] for r in rr];med=statistics.median(vals);mad=statistics.median(abs(v-med) for v in vals)
   for r in rr:
    value=r[k];score=.67448975*(value-med)/mad if mad else 0
    if abs(score)>3.5 and abs(value-med)>=minimum:
     peers=[a[k] for a in rr if a is not r]
     flags.append(dict(sample_id=r['Probenr-HAFL'],excel_row=r['excel_row'],barn=barn,mastday=day,metric=k,value=value,peer_min=min(peers),peer_max=max(peers),peer_mean=statistics.mean(peers),group_median=med,robust_z=score))
def save_csv(name,rows):
 with (out/name).open('w',encoding='utf-8-sig',newline='') as f:
  w=csv.DictWriter(f,fieldnames=list(rows[0]),delimiter=';');w.writeheader();w.writerows(rows)
save_csv('Plausibilitaet_August2026_Gruppen.csv',groups)
save_csv('Plausibilitaet_August2026_Rechenkontrollen.csv',checks)
save_csv('Plausibilitaet_August2026_Prueffaelle.csv',flags)
save_csv('Plausibilitaet_August2026_Kampagnenvergleich.csv',ranges)
def fmt(x,d=2):return f'{x:.{d}f}'.replace('.',',')
def link(p):return f'[{p.name}](<{p}>)'
report=['# Plausibilitätsprüfung Einstreudaten August 2026','',
 'Stand: 29.10.2026. Prüfung ohne Änderung der Gesamtdatei oder Laborwerte.','',
 '**Ergebnis:** Die 31 August-Proben sind korrekt zugeordnet und quellengleich übertragen. Die Messwerte zeigen keine groben physikalischen Widersprüche. Vor einer Nährstoffauswertung müssen aber bestehende Einheiten- und Phosphorbezeichnungen bereinigt werden. Einzelne Laborwerte sind innerhalb ihrer Probenahmegruppe auffällig und sollten geprüft werden.','',
 '## Quellen und Umfang','',
 '- Gesamtdatei: '+link(target)+', Blatt Sheet 1, Zeilen 154–184, Kampagne AUG26.',
 '- Probenahme: '+link(meta_path)+', Blatt Sheet1, A2:E32.',
 '- Labor: '+link(lab_path)+', Blatt 125956_BFH_Gfeller_20260917, G5:AG35; Parameter und Einheiten in Zeilen 1–3.',
 '- Mastbeginn aus local-times.rds: beide Ställe am 24.07.2026.',
 f'- {n_source} Zuordnungs- und Messwertzellen mit den Quellen abgeglichen; keine Abweichung. Feuchte und N in Frischsubstanz zusätzlich nachgerechnet.',
 '- Alle 31 Probennummern besitzen Laborbefund und Probenahmezuordnung. Keine doppelten Kombinationen aus Kampagne, Stall, Masttag und Wiederholung.',
 '- 13.08., 18.08. und 26.08.2026 entsprechen den Masttagen 21, 26 und 34.',
 '- Stall 1 hat 16, Stall 2 hat 15 Proben. Die zusätzliche Probe 620 ist als Wiederholung 6, Stall 1, Masttag 34 dokumentiert.','',
 '## Augustwerte je Stall und Masttag','',
 'Arithmetische Mittel der Proben. Die Proben sind keine 31 unabhängigen Ställe oder Kampagnen. Beide Ställe liefen ohne WRG; aus dem Vergleich wird kein WRG-Effekt abgeleitet.','',
 '| Stall | Masttag | n | pH | Feuchte [%] | N gesamt [g/kg TS] | NH4-N [g/kg TS] | C/N |',
 '|---|---:|---:|---:|---:|---:|---:|---:|']
for g in groups:
 report.append('| '+' | '.join([str(g['barn']),str(g['mastday']),str(g['n'])]+[fmt(g[k+'_mean']) for k in metrics])+' |')
report+=['',
 'NH4-N folgt dem Laborparameter «Ammoniumstickstoff». Die Kurzbezeichnung NH4_gkg_TS sollte nicht als Masse des gesamten NH4-Ions interpretiert werden.','',
 'Stall 1: Feuchte und N steigen mit dem Mastalter, das C/N-Verhältnis sinkt. Stall 2: Am Tag 34 ist die Einstreu mit 26,02 % mittlerer Feuchte trockener als am Tag 26 (31,00 %) und als Stall 1 am Tag 34 (36,38 %). Der Stallunterschied beträgt am Tag 34 10,36 Prozentpunkte. Alle fünf Tag-34-Proben aus Stall 2 liegen zwischen 25,3 und 27,2 % Feuchte; die Beobachtung hängt nicht an einem Einzelwert. Eine Ursache lässt sich daraus ohne Probenpositionen und weitere Prozessdaten nicht bestimmen.','',
 '## Vergleich mit früheren Kampagnen','',
 '| Parameter | August 2026: Minimum–Maximum | Frühere Kampagnen: Minimum–Maximum |',
 '|---|---:|---:|']
for k in metrics+['TS_pct']:
 a=[r[k] for r in aug];o=[r[k] for r in old]
 report.append(f'| {k} | {fmt(min(a))}–{fmt(max(a))} | {fmt(min(o))}–{fmt(max(o))} |')
report+=['',
 'Historische Spannweiten dienen als interne Orientierung, nicht als validierte Grenzwerte. Kampagnen besitzen unterschiedliche Masttage und Probenzahlen. Im Juli 2024 steht bereits pH 0,52 für Probe 321 (Gesamtdatei F102). Dieser auffällige Wert steht auch in der ursprünglichen Juli-Labordatei P26 und in deren editierter Fassung P22. Er ist kein August-Importfehler und sollte nicht als plausibles Referenzminimum verwendet werden.',
 '- Historische Originalquelle: '+link(Path('/Users/stefan/git/mettmenstetten-jul2024/data/HD_116922_HAFL Gfeller Stefan_20241106.xlsx'))+'.','',
 '## Rechenkontrollen','',
 '- Feuchte = 100 − TS und N_FS = N_TS × TS/100: jeweils alle 31 Werte korrekt.',
 '- N_FS laut Labor weicht maximal 0,0508 g/kg FS vom Ergebnis aus angezeigten gerundeten Eingangswerten ab; alle Abweichungen passen zur angegebenen Stellenzahl.',
 '- C/N = Corg/N_TS: maximale Abweichung 0,016; alle Werte mit der Rundung vereinbar.',
 '- Glühverlust + Asche = 100,0 bis 100,1 %. Kleine Rundungsdifferenz, kein Übertragungsfehler.',
 '- P2O5/P ≈ 2,291 und K2O/K ≈ 1,205: alle Werte innerhalb der Rundungstoleranzen der Eingangswerte.',
 '- NH4-N liegt bei jeder Probe zwischen 0 und Gesamt-N. Keine negativen Gehalte, keine TS- oder Feuchtewerte ausserhalb 0–100 %.',
 '- Corg liegt unter der Trockenmasse; die Summe der Mineralstoffelemente liegt unter dem Aschegehalt, wenn die Laboreinheiten kg/t TS bzw. g/kg TS verwendet werden.',
 '- Corg entspricht weitgehend einer Berechnung aus organischem Anteil mit Faktor 0,58 und ist damit keine unabhängige Bestätigung des Glühverlustes. Probe 618: 87,6 % Glühverlust plus 12,5 % Asche; Corg 507,5 g/kg entspricht (100 − 12,5) × 5,8. Diese kleine Uneinheitlichkeit ist bereits im Laborbefund enthalten.','',
 '## Wichtige Bezeichnungsprobleme','',
 '**Mineralstoffe N:T: kg/t TS statt g/t TS.** Die Überschriften nennen g/t TS, das Labor weist kg/t TS aus. Die Zahlen folgen unverändert der bisherigen Skala. Eine Interpretation als g/t wäre um Faktor 1000 falsch. Sinnvolle Bereinigung: Überschriften auf kg/t TS oder g/kg TS korrigieren, nicht die Zahlen mit 1000 multiplizieren.','',
 '**Spalte V enthält offenbar P2O5_FS statt elementarem P_FS.** Beispiel Probe 431, Zeile 154: 4,77 kg P/t TS × 78,6 % TS = 3,74922 g P/kg FS. V154 enthält 8,59 g/kg, nahe bei 10,94 × 0,786 = 8,59884 g P2O5/kg FS. Diese Bezeichnung stammt schon aus dem Laborfeld AC und tritt auch in den älteren Kampagnen auf. Spalte W ist als Phosphorpentoxid g/kg FS beschriftet, aber leer.','',
 'Mögliche spätere Bereinigung: bisherige V-Werte als Labor-P2O5_FS in W erhalten und elementaren P_FS in V aus Q × G/100 berechnen. Die Umrechnung P → P2O5 mit Faktor 2,291 ist in [BAFU/Agroscope: Le phosphore dans les sols, Seite 113](https://www.bafu.admin.ch/dam/fr/sd-web/aAP0-NeynzHD/phosphor_in_boedenstandortbestimmungschweizphosphorinboedendueng.pdf#page=114) dokumentiert.','',
 '## Auffällige Einzelproben','',
 'Screening innerhalb gleicher Ställe und Masttage: |modifizierter z-Wert| > 3,5 sowie Mindestabweichung vom Median: pH 0,3; TS 4 Prozentpunkte; N 4 g/kg TS; NH4-N 0,8 g/kg TS. Die Schwellen dienen der Priorisierung. Kleine Gruppen und kleine MAD können grosse z-Werte erzeugen; kein Hinweis beweist einen Messfehler. Alle aufgeführten Werte sind im Laborbefund identisch enthalten.','',
 '| Probe | Excelzeile | Stall/Tag | Merkmal | Wert | Andere Proben der Gruppe |',
 '|---|---:|---|---|---:|---:|']
for f in flags:
 report.append(f'| {f["sample_id"]} | {f["excel_row"]} | {f["barn"]}/{f["mastday"]} | {f["metric"]} | {fmt(f["value"])} | {fmt(f["peer_min"])}–{fmt(f["peer_max"])} |')
report+=['',
 'Besonders prüfenswert: Probe 438 mit pH 6,57 gegenüber 5,63–5,91 in den anderen Tag-21-Proben aus Stall 2, Probe 606 mit pH 6,01 gegenüber 6,36–6,46 an Tag 26 sowie Probe 610 mit NH4-N 2,10 gegenüber 3,61–3,90 g/kg TS an Tag 34. Probe 603 ist relativ feucht und hat zugleich höheres N pro TS; diese Kombination kann durch die andere Bezugsmasse verstärkt werden.','',
 'Sensitivität: Ohne Probe 610 wäre der NH4-N-Mittelwert von Stall 2 an Tag 34 3,7075 statt 3,386 g/kg TS (+9,5 %). Ohne Probe 438 wäre der pH-Mittelwert von Stall 2 an Tag 21 5,8075 statt 5,96. Die pH-Mittel sind nur deskriptiv, keine Berechnung der Aktivität einer Mischprobe. Es wurden keine Werte entfernt oder geändert.','',
 '## Dateien','',
 '- Plausibilitaet_August2026_Gruppen.csv: Gruppengrössen, Mittel, Streuung und Spannweiten.',
 '- Plausibilitaet_August2026_Rechenkontrollen.csv: Einzelprüfungen, Residuen und Rundungstoleranzen.',
 '- Plausibilitaet_August2026_Prueffaelle.csv: priorisierte Auffälligkeiten.',
 '- Plausibilitaet_August2026_Kampagnenvergleich.csv: deskriptive Kampagnenkennzahlen.','',
 'Die Prüfung bestätigt Zuordnung und Rechenkonsistenz, nicht die analytische Richtigkeit der Laborbestimmungen. Für auffällige Einzelproben wären Labor-Rohprotokolle oder Wiederholungsbestimmungen erforderlich.']
report+=['','## Präzisierung zum Probenahmeverdacht','',
 'Der Hinweis bezieht sich auf den Mittelwert der Feuchtigkeit aller fünf Proben aus Stall 2 am Masttag 34 (26,02 %), nicht auf den NH4-N-Einzelwert der Probe 610. Vermutet wird eine für die Einstreu im Stall nicht repräsentative Probenahme. Die Messwerte bleiben unverändert; der Verdacht wurde nicht durch eine Wiederholungsbeprobung bestätigt.']
(out/'Plausibilitaet_August2026.md').write_text('\n'.join(report)+'\n')
assert hashlib.sha256(target.read_bytes()).hexdigest()==sha
print(json.dumps(dict(source_cells_verified=n_source,arithmetic_checks=len(checks),arithmetic_passes=sum(c['within_rounding'] for c in checks),prioritized_flags=flags,report=str(out/'Plausibilitaet_August2026.md'),workbook_unchanged=True),ensure_ascii=False,indent=2))
